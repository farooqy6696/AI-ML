"""
Report Module (Service Layer)
Generates Excel reports, PDF reports, department-wise performance reports,
skill distribution charts, and an employee summary dashboard.
"""
import os
import uuid
from collections import Counter, defaultdict
from typing import List

import pandas as pd
from openpyxl.chart import BarChart, Reference
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer

from app.config import settings
from app.models.employee import Employee


def _employees_to_dataframe(employees: List[Employee]) -> pd.DataFrame:
    rows = [
        {
            "Employee ID": e.employee_id,
            "Name": e.name,
            "Email": e.email,
            "Department": e.department,
            "Designation": e.designation,
            "Experience (yrs)": e.experience,
            "Skills": e.skills,
            "Certifications": e.certifications,
            "Performance Score": e.performance_score,
            "Attendance (%)": e.attendance,
            "Salary": e.salary,
            "Projects Completed": e.projects_completed,
            "Manager Feedback": e.manager_feedback,
        }
        for e in employees
    ]
    return pd.DataFrame(rows)


class ReportService:
    def __init__(self):
        os.makedirs(settings.REPORTS_DIR, exist_ok=True)

    def _output_path(self, prefix: str, ext: str) -> str:
        return os.path.join(settings.REPORTS_DIR, f"{prefix}_{uuid.uuid4().hex[:8]}.{ext}")

    # ---- Excel Report ----
    def generate_excel_report(self, employees: List[Employee]) -> str:
        df = _employees_to_dataframe(employees)
        path = self._output_path("employee_report", "xlsx")

        with pd.ExcelWriter(path, engine="openpyxl") as writer:
            df.to_excel(writer, index=False, sheet_name="Employees")

            dept_summary = df.groupby("Department").agg(
                Avg_Performance=("Performance Score", "mean"),
                Avg_Attendance=("Attendance (%)", "mean"),
                Headcount=("Employee ID", "count"),
            ).reset_index()
            dept_summary.to_excel(writer, index=False, sheet_name="Department Summary")

            worksheet = writer.sheets["Employees"]
            for i, col in enumerate(df.columns, start=1):
                max_len = max(df[col].astype(str).map(len).max() if len(df) else 0, len(col)) + 2
                worksheet.column_dimensions[get_column_letter(i)].width = min(max_len, 40)

            # Bar chart: average performance by department
            chart_sheet = writer.sheets["Department Summary"]
            chart = BarChart()
            chart.title = "Average Performance by Department"
            chart.y_axis.title = "Avg Performance Score"
            chart.x_axis.title = "Department"
            data = Reference(chart_sheet, min_col=2, min_row=1, max_row=len(dept_summary) + 1)
            cats = Reference(chart_sheet, min_col=1, min_row=2, max_row=len(dept_summary) + 1)
            chart.add_data(data, titles_from_data=True)
            chart.set_categories(cats)
            chart_sheet.add_chart(chart, "G2")

        return path

    # ---- PDF Report ----
    def generate_pdf_report(self, employees: List[Employee]) -> str:
        path = self._output_path("employee_report", "pdf")
        doc = SimpleDocTemplate(path, pagesize=A4)
        styles = getSampleStyleSheet()
        elements = [Paragraph("Employee Performance Report", styles["Title"]), Spacer(1, 12)]

        table_data = [
            ["ID", "Name", "Department", "Designation", "Performance", "Attendance", "Salary"]
        ]
        for e in employees:
            table_data.append(
                [
                    str(e.employee_id),
                    e.name,
                    e.department,
                    e.designation,
                    f"{e.performance_score or 0:.1f}",
                    f"{e.attendance or 0:.1f}%",
                    f"{e.salary or 0:,.0f}",
                ]
            )

        table = Table(table_data, repeatRows=1)
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2c3e50")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("FONTSIZE", (0, 0), (-1, -1), 8),
                    ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f2f2f2")]),
                ]
            )
        )
        elements.append(table)
        doc.build(elements)
        return path

    # ---- Department-wise Performance Report ----
    def department_wise_report(self, employees: List[Employee]) -> dict:
        df = _employees_to_dataframe(employees)
        if df.empty:
            return {}
        summary = df.groupby("Department").agg(
            avg_performance=("Performance Score", "mean"),
            avg_attendance=("Attendance (%)", "mean"),
            avg_salary=("Salary", "mean"),
            headcount=("Employee ID", "count"),
        ).round(2)
        return summary.to_dict(orient="index")

    # ---- Skill Distribution Charts (data) ----
    def skill_distribution(self, employees: List[Employee]) -> dict:
        counter = Counter()
        for e in employees:
            for skill in e.skills_list():
                counter[skill] += 1
        return dict(counter.most_common(20))

    # ---- Employee Summary Dashboard (data) ----
    def summary_dashboard(self, employees: List[Employee]) -> dict:
        if not employees:
            return {
                "total_employees": 0,
                "avg_performance": 0,
                "avg_attendance": 0,
                "avg_salary": 0,
                "department_counts": {},
                "top_performers": [],
            }
        df = _employees_to_dataframe(employees)
        dept_counts = df["Department"].value_counts().to_dict()
        top_performers = (
            df.sort_values("Performance Score", ascending=False)
            .head(5)[["Name", "Department", "Performance Score"]]
            .to_dict(orient="records")
        )
        return {
            "total_employees": len(employees),
            "avg_performance": round(df["Performance Score"].mean(), 2),
            "avg_attendance": round(df["Attendance (%)"].mean(), 2),
            "avg_salary": round(df["Salary"].mean(), 2),
            "department_counts": dept_counts,
            "top_performers": top_performers,
        }


report_service = ReportService()
